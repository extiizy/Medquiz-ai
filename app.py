import os, json, re, tempfile
from pathlib import Path
from flask import Flask, request, jsonify, send_from_directory
from openai import OpenAI
from pptx import Presentation
from pypdf import PdfReader

BASE = Path(__file__).resolve().parent
app = Flask(__name__, static_folder="static")
app.config["MAX_CONTENT_LENGTH"] = 30 * 1024 * 1024

def extract_pptx(path):
    prs = Presentation(path)
    chunks = []
    for si, slide in enumerate(prs.slides, 1):
        texts = []
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                texts.append(shape.text.strip())
        if texts:
            chunks.append(f"[Slide {si}]\n" + "\n".join(texts))
    return "\n\n".join(chunks)

def extract_pdf(path):
    reader = PdfReader(path)
    chunks = []
    for i, page in enumerate(reader.pages, 1):
        txt = page.extract_text() or ""
        if txt.strip():
            chunks.append(f"[Page {i}]\n{txt.strip()}")
    return "\n\n".join(chunks)

def extract_material(path):
    ext = path.suffix.lower()
    if ext == ".pptx":
        return extract_pptx(path)
    if ext == ".pdf":
        return extract_pdf(path)
    raise ValueError("Format belum didukung. Gunakan PPTX atau PDF.")

def get_client():
    key = os.getenv("OPENAI_API_KEY")
    if not key:
        raise RuntimeError("OPENAI_API_KEY belum diatur di server.")
    return OpenAI(api_key=key)

def generate_ai(material, count=10, difficulty="sedang"):
    # Keep the source as the primary authority; the model must not silently invent
    # unsupported facts. If a requested point is absent, it must say so.
    prompt = f"""
Kamu adalah mesin belajar kedokteran untuk mahasiswa kedokteran Indonesia.
Buat soal latihan bergaya UKMPPD berdasarkan MATERI SUMBER di bawah.

ATURAN SUMBER:
- Materi sumber adalah sumber utama.
- Pertahankan istilah, klasifikasi, angka, dan framing dari sumber.
- Jangan mengarang fakta seolah-olah berasal dari sumber.
- Jika informasi yang dibutuhkan untuk sebuah soal tidak ada di sumber, jangan pakai
  informasi tersebut sebagai fakta utama; tandai sebagai "informasi tambahan" jika perlu.
- Ini adalah latihan pendidikan, bukan pengganti keputusan klinis.

TARGET:
- Jumlah soal: {count}
- Kesulitan: {difficulty}
- Bahasa: Indonesia.
- Buat vignette klinis yang realistis, tetapi tetap dapat dijawab dari materi.
- Setiap soal punya tepat 5 opsi A-E dan satu jawaban terbaik.
- Pembahasan harus menjelaskan mengapa jawaban benar dan mengapa masing-masing opsi lain salah.
- Tambahkan "high_yield" dan "flashcards".
- Flashcard ringkas tetapi bermakna.

Kembalikan HANYA JSON valid dengan struktur:
{{
  "title": "judul materi",
  "source_note": "ringkasan singkat sumber",
  "questions": [
    {{
      "question": "...",
      "options": {{"A":"...","B":"...","C":"...","D":"...","E":"..."}},
      "answer": "A",
      "explanation": "...",
      "option_explanations": {{"A":"...","B":"...","C":"...","D":"...","E":"..."}},
      "high_yield": "..."
    }}
  ],
  "flashcards": [
    {{"front":"...","back":"..."}}
  ]
}}

MATERI SUMBER:
{material[:180000]}
"""
    client = get_client()
    response = client.responses.create(
        model=os.getenv("OPENAI_MODEL", "gpt-6-luna"),
        input=prompt
    )
    raw = response.output_text.strip()
    raw = re.sub(r"^```json\s*", "", raw)
    raw = re.sub(r"\s*```$", "", raw)
    return json.loads(raw)

@app.get("/")
def index():
    return send_from_directory(app.static_folder, "index.html")

@app.post("/api/generate")
def generate():
    if "file" not in request.files:
        return jsonify(error="File belum dipilih."), 400
    f = request.files["file"]
    if not f.filename:
        return jsonify(error="Nama file kosong."), 400
    suffix = Path(f.filename).suffix.lower()
    if suffix not in {".pptx", ".pdf"}:
        return jsonify(error="Gunakan file PPTX atau PDF."), 400
    try:
        count = max(1, min(int(request.form.get("count", 10)), 20))
        difficulty = request.form.get("difficulty", "sedang")
    except ValueError:
        count, difficulty = 10, "sedang"

    with tempfile.NamedTemporaryFile(delete=False, suffix=suffix) as tmp:
        f.save(tmp.name)
        temp_path = Path(tmp.name)
    try:
        material = extract_material(temp_path)
        if not material.strip():
            return jsonify(error="Tidak ditemukan teks pada file."), 400
        result = generate_ai(material, count, difficulty)
        return jsonify(result)
    except Exception as e:
        return jsonify(error=str(e)), 500
    finally:
        try: temp_path.unlink()
        except Exception: pass

if __name__ == "__main__":
    app.run(host="0.0.0.0", port=int(os.getenv("PORT", "8000")))
